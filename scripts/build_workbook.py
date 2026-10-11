"""Build the benchmark workbook (results/brightline_benchmark.xlsx) from the run logs.

Sheets
    Board           one row per configuration, live formulas over the Runs sheet, budget check,
                    trade-off charts, and the rank/verdict cells to fill in
    Runs            every benchmark trial from results/runs.csv (raw data, nothing edited)
    Training        parameter counts, training tokens, time, validation accuracy per config
    Scaling         6ND training-compute estimate and Chinchilla token budget for a 10x model
    Recommendation  ranked recommendation + residual-risk register (decision cells)
    Protocol        how every number was measured, environment, dataset composition

The decision cells (Board Rank/Verdict, Recommendation rows) are filled from
results/recommendation.json when it exists; edit that file rather than the workbook so a
rebuild never wipes the ranking.

Board and Scaling are formulas, not pasted numbers: re-run benchmarks, re-run this script,
and every figure updates. The Board reads the latest run of each config (Run ID column);
type an older run_id there to pin a different run.

Usage:
    python scripts/build_workbook.py
"""
from __future__ import annotations

import csv
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference, ScatterChart, Series
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import DATA_DIR, RESULTS_DIR, RUN_FIELDS, RUNS_CSV, WORKBOOK  # noqa: E402

RECOMMENDATION = RESULTS_DIR / "recommendation.json"
from configs import CONFIGS  # noqa: E402

FONT = "Arial"
BLUE, GREEN, BLACK, GREY = "0000FF", "008000", "000000", "666666"
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
BAND_FILL = PatternFill("solid", fgColor="D9E1F2")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NAMES = list(CONFIGS)
SERIES_COLORS = ["2F5597", "ED7D31", "70AD47", "7030A0", "C00000", "00B0F0"]  # one per config, in order


def font(color=BLACK, bold=False, size=10, italic=False):
    return Font(name=FONT, color=color, bold=bold, size=size, italic=italic)


def header_row(ws, row, labels, col=1):
    for i, label in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=label)
        c.font = Font(name=FONT, color="FFFFFF", bold=True, size=10)
        c.fill = HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = BOX


def title(ws, text, subtitle=None):
    ws["A1"] = text
    ws["A1"].font = font(bold=True, size=14)
    if subtitle:
        ws["A2"] = subtitle
        ws["A2"].font = font(GREY, italic=True)


def put(ws, ref, value, color=BLACK, fmt=None, bold=False, fill=None, note=None):
    c = ws[ref]
    c.value = value
    c.font = font(color, bold=bold)
    c.border = BOX
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if note:
        c.comment = Comment(note, "benchmark")
    return c


def set_widths(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


# --- data loading ------------------------------------------------------------------------
def load_runs() -> list[dict]:
    if not RUNS_CSV.exists():
        return []
    with open(RUNS_CSV, encoding="utf-8") as f:
        return list(csv.DictReader(f))


def latest_run_ids(runs: list[dict]) -> dict:
    latest = {}
    for r in runs:  # runs.csv is append-only, so the last run_id seen per config is the newest
        latest[r["config"]] = r["run_id"]
    return latest


def load_training() -> dict:
    out = {}
    for name in NAMES:
        p = RESULTS_DIR / f"train_{name}.json"
        if p.exists():
            out[name] = json.loads(p.read_text())
    return out


def numeric(field: str, value: str):
    if value == "":
        return None
    if field in {"run_id", "timestamp_utc", "config", "variant_type", "description", "precision", "device",
                 "torch_version", "cuda_version", "python_version"}:
        return value
    try:
        return int(value) if value.lstrip("-").isdigit() else float(value)
    except ValueError:
        return value


# --- sheets --------------------------------------------------------------------------------
def build_runs(wb, runs):
    ws = wb.create_sheet("Runs")
    header_row(ws, 1, RUN_FIELDS)
    for i, r in enumerate(runs, start=2):
        for j, field in enumerate(RUN_FIELDS, start=1):
            c = ws.cell(row=i, column=j, value=numeric(field, r.get(field, "")))
            c.font = font()
    last = max(2, len(runs) + 1)
    if runs:
        ws.add_table(Table(displayName="RunsTable", ref=f"A1:{get_column_letter(len(RUN_FIELDS))}{last}",
                           tableStyleInfo=TableStyleInfo(name="TableStyleLight9", showRowStripes=True)))
    ws.freeze_panes = "D2"
    for j, field in enumerate(RUN_FIELDS, start=1):
        ws.column_dimensions[get_column_letter(j)].width = max(11, min(40, len(field) + 3))
    ws.column_dimensions["E"].width = 48
    return last


def first_run(field: str) -> str:
    """Formula fragment: value of `field` in the first Runs data row (environment is the same for all rows)."""
    col = get_column_letter(RUN_FIELDS.index(field) + 1)
    return f"INDEX(Runs!${col}:${col},2)"


def runs_col(field: str, last: int) -> str:
    col = get_column_letter(RUN_FIELDS.index(field) + 1)
    return f"Runs!${col}$2:${col}${last}"


def build_board(wb, runs, last, ranking):
    ws = wb.active
    ws.title = "Board"
    title(ws, "Brightline Logistics - In-Cab Model Trade-off Board",
          "Scanned bills of lading + spoken driver notes, one 16-class model. "
          "Every figure is the mean of repeated benchmark trials (see Protocol).")
    ws["A3"] = (f'=IF(COUNTA(Runs!$A:$A)<2,"Measured on: (no runs yet)","Measured on: "&{first_run("device")}'
                f'&"  |  torch "&{first_run("torch_version")})')
    ws["A3"].font = font(GREEN)
    ws["A4"] = ("Legend:  blue = input you can change  |  yellow fill = decision cells to fill in  |  "
                "green = pulled from another sheet  |  black = formula")
    ws["A4"].font = font(GREY, italic=True)

    # Budget thresholds (assumptions Sandra's team should set).
    ws["A6"] = "Tablet budget (edit to match the agreed hardware target)"
    ws["A6"].font = font(bold=True)
    budgets = [
        ("Accuracy floor (overall)", 0.85, "0.0%", "Assumption: minimum acceptable overall accuracy. Set with ops."),
        ("p95 latency ceiling (ms, measured GPU)", 10, "0.0", "Assumption: p95 per-request latency budget on "
         "the benchmark GPU. The tablet will be slower; see Recommendation > risks."),
        ("Shipped model size ceiling (MB)", 8, "0.0", "Assumption: max serialized weight size per OTA update."),
    ]
    for i, (label, val, fmt, note) in enumerate(budgets, start=7):
        put(ws, f"A{i}", label)
        put(ws, f"B{i}", val, BLUE, fmt, fill=INPUT_FILL, note=note)

    hdr = 11
    cols = ["Config", "Run ID", "Variant", "What changed", "Trials", "Accuracy", "Docs acc", "Notes acc",
            "p50 ms", "p95 ms", "p50 spread across trials (ms)", "Peak GPU mem (MB)", "Shipped size (MB)",
            "Params", "Accuracy vs baseline (pts)", "p95 speed-up vs baseline (x)", "Size reduction vs baseline (x)",
            "Meets budget?", "Rank", "Verdict"]
    header_row(ws, hdr, cols)
    ws.row_dimensions[hdr].height = 42

    latest = latest_run_ids(runs)
    first = hdr + 1
    rid = lambda r: f"$B{r}"  # noqa: E731
    for i, name in enumerate(NAMES):
        r = first + i
        cfg = CONFIGS[name]
        put(ws, f"A{r}", name, bold=True)
        put(ws, f"B{r}", latest.get(name, ""), BLUE,
            note="run_id from the Runs sheet. Defaults to the newest run of this config." if i == 0 else None)
        put(ws, f"C{r}", cfg["variant_type"])
        put(ws, f"D{r}", cfg["description"])

        def avg(field):
            return f"=IFERROR(AVERAGEIFS({runs_col(field, last)},{runs_col('run_id', last)},{rid(r)}),\"\")"

        put(ws, f"E{r}", f"=IF({rid(r)}=\"\",\"\",COUNTIFS({runs_col('run_id', last)},{rid(r)}))", GREEN, "0")
        put(ws, f"F{r}", avg("accuracy"), GREEN, "0.0%")
        put(ws, f"G{r}", avg("acc_synthetic_docs"), GREEN, "0.0%")
        put(ws, f"H{r}", avg("acc_slurp_notes"), GREEN, "0.0%")
        put(ws, f"I{r}", avg("p50_ms"), GREEN, "0.00")
        put(ws, f"J{r}", avg("p95_ms"), GREEN, "0.00")
        p50, ids = runs_col("p50_ms", last), runs_col("run_id", last)
        put(ws, f"K{r}", f"=IFERROR(IF(E{r}=\"\",\"\",_xlfn.MAXIFS({p50},{ids},{rid(r)})-_xlfn.MINIFS({p50},{ids},"
                         f"{rid(r)})),\"\")", GREEN, "0.00")
        put(ws, f"L{r}", avg("peak_mem_mb"), GREEN, "0.0")
        put(ws, f"M{r}", avg("checkpoint_mb"), GREEN, "0.00")
        put(ws, f"N{r}", avg("params_total"), GREEN, "#,##0")
        put(ws, f"O{r}", f"=IFERROR((F{r}-F${first})*100,\"\")", fmt='+0.0;-0.0;0.0')
        put(ws, f"P{r}", f"=IFERROR(J${first}/J{r},\"\")", fmt="0.00\"x\"")
        put(ws, f"Q{r}", f"=IFERROR(M${first}/M{r},\"\")", fmt="0.0\"x\"")
        put(ws, f"R{r}", f"=IF(F{r}=\"\",\"\",IF(AND(F{r}>=$B$7,J{r}<=$B$8,M{r}<=$B$9),\"PASS\",\"FAIL\"))")
        rank = next((k for k, e in enumerate(ranking, start=1) if e["config"] == name), None)
        put(ws, f"S{r}", rank, BLUE, "0", fill=INPUT_FILL)
        put(ws, f"T{r}", ranking[rank - 1]["verdict"] if rank else None, BLUE, fill=INPUT_FILL)
        ws[f"T{r}"].alignment = Alignment(wrap_text=True, vertical="top")
        ws[f"D{r}"].alignment = Alignment(wrap_text=True, vertical="top")
    lastrow = first + len(NAMES) - 1
    ws[f"S{first}"].comment = Comment("Fill in 1 = ship this. Ranking rationale goes on the Recommendation sheet.",
                                      "benchmark")

    dv = DataValidation(type="whole", operator="between", formula1="1", formula2=str(len(NAMES)), allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"S{first}:S{lastrow}")
    ws.conditional_formatting.add(f"R{first}:R{lastrow}", CellIsRule(operator="equal", formula=['"PASS"'],
                                  fill=PatternFill("solid", fgColor="C6EFCE"), font=Font(name=FONT, color="006100")))
    ws.conditional_formatting.add(f"R{first}:R{lastrow}", CellIsRule(operator="equal", formula=['"FAIL"'],
                                  fill=PatternFill("solid", fgColor="FFC7CE"), font=Font(name=FONT, color="9C0006")))
    for col in "FGH":
        ws.conditional_formatting.add(f"{col}{first}:{col}{lastrow}", ColorScaleRule(
            start_type="min", start_color="F8CBAD", end_type="max", end_color="C6EFCE"))
    for col in "IJLM":
        ws.conditional_formatting.add(f"{col}{first}:{col}{lastrow}", ColorScaleRule(
            start_type="min", start_color="C6EFCE", end_type="max", end_color="F8CBAD"))

    note_row = lastrow + 2
    ws[f"A{note_row}"] = ("Accuracy = 16-class test accuracy (1,600 unseen-customer documents + 893 SLURP notes). "
                          "Latency = one request at a time, host->GPU->host, synchronized. "
                          "Shipped size = serialized weights as loaded (4-bit packs two weights per byte).")
    ws[f"A{note_row}"].font = font(GREY, italic=True)

    # Charts: accuracy vs p95 latency (one series per config so each is labelled) + shipped size.
    sc = ScatterChart()
    sc.title = "Accuracy vs p95 latency (up and left is better)"
    sc.style = 13
    sc.x_axis.title = "p95 latency (ms)"
    sc.y_axis.title = "Accuracy"
    sc.y_axis.number_format = "0.0%"
    sc.x_axis.number_format = "0.0"
    sc.x_axis.delete = False
    sc.y_axis.delete = False
    sc.x_axis.majorGridlines = None
    sc.legend.position = "r"
    sc.legend.overlay = False
    sc.title.overlay = False
    for i, r in enumerate(range(first, lastrow + 1)):
        s = Series(Reference(ws, min_col=6, min_row=r), Reference(ws, min_col=10, min_row=r),
                   title_from_data=False, title=ws[f"A{r}"].value)
        color = SERIES_COLORS[i % len(SERIES_COLORS)]
        s.marker.symbol = "circle"
        s.marker.size = 11
        s.marker.graphicalProperties.solidFill = color
        s.marker.graphicalProperties.line.solidFill = color
        s.graphicalProperties.line.noFill = True
        sc.series.append(s)
    sc.height, sc.width = 9, 18
    ws.add_chart(sc, f"A{note_row + 2}")

    bc = BarChart()
    bc.type = "bar"
    bc.title = "Shipped model size (MB)"
    bc.style = 10
    bc.add_data(Reference(ws, min_col=13, min_row=hdr, max_row=lastrow), titles_from_data=True)
    bc.set_categories(Reference(ws, min_col=1, min_row=first, max_row=lastrow))
    bc.legend = None
    bc.title.overlay = False
    bc.series[0].graphicalProperties.solidFill = SERIES_COLORS[0]
    bc.x_axis.majorGridlines = None
    bc.x_axis.delete = False
    bc.y_axis.delete = False
    bc.height, bc.width = 9, 14
    ws.add_chart(bc, f"I{note_row + 2}")

    set_widths(ws, {"A": 34, "B": 11, "C": 13, "D": 46, "E": 8, "F": 10, "G": 10, "H": 10, "I": 9, "J": 9,
                    "K": 12, "L": 11, "M": 11, "N": 12, "O": 12, "P": 12, "Q": 12, "R": 10, "S": 7, "T": 48})
    ws.freeze_panes = f"B{first}"
    return first, lastrow


def build_training(wb, training):
    ws = wb.create_sheet("Training")
    cols = ["config", "variant_type", "teacher", "params_total", "params_non_embedding", "train_tokens", "epochs",
            "train_examples_per_epoch", "best_val_acc", "train_time_s", "device"]
    header_row(ws, 1, cols)
    fmts = {"params_total": "#,##0", "params_non_embedding": "#,##0", "train_tokens": "#,##0",
            "train_examples_per_epoch": "#,##0", "best_val_acc": "0.0%", "train_time_s": "0.0"}
    for i, name in enumerate(NAMES, start=2):
        cfg, t = CONFIGS[name], training.get(name)
        values = {"config": name, "variant_type": cfg["variant_type"], "teacher": cfg.get("teacher", "")}
        if "source" in cfg:
            values["device"] = f"post-training transform of {cfg['source']} - no training run"
        elif t:
            values.update({k: t.get(k) for k in cols[3:]})
        else:
            values["device"] = "(not trained yet)"
        for j, k in enumerate(cols, start=1):
            c = ws.cell(row=i, column=j, value=values.get(k))
            c.font = font()
            c.border = BOX
            if k in fmts:
                c.number_format = fmts[k]
    n = len(NAMES) + 3
    ws.cell(row=n, column=1, value="train_tokens = real (non-padding) bytes the model processed over all epochs: "
            "the D in 6ND. The distilled student also ran the teacher forward on every batch (~2*N_teacher*D extra "
            "FLOPs), which is not included in its own 6ND.").font = font(GREY, italic=True)
    set_widths(ws, {get_column_letter(j): 16 for j in range(1, len(cols) + 1)})
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["K"].width = 44


def build_scaling(wb, avg_tokens):
    ws = wb.create_sheet("Scaling")
    title(ws, "Scaling-law compute estimate: what would a 10x larger model cost?",
          "Training FLOPs ~= 6 x N (parameters) x D (training tokens).  Chinchilla compute-optimal budget: "
          "D ~= 20 tokens per parameter (Hoffmann et al., 2022).")
    put(ws, "A4", "Inputs", bold=True)
    tr = "INDEX(Training!${col}:${col},MATCH(\"baseline\",Training!$A:$A,0))"
    inputs = [
        ("Baseline parameters N", "=" + tr.format(col="D"), GREEN, "#,##0", "From Training sheet (total params, "
         "Chinchilla convention)."),
        ("Baseline training tokens D (as trained)", "=" + tr.format(col="F"), GREEN, "#,##0", "From Training sheet."),
        ("Unique training tokens per epoch", "=" + tr.format(col="F") + "/" + tr.format(col="G"), GREEN, "#,##0",
         "D / epochs: how much distinct data we actually have."),
        ("Scale factor", 10, BLUE, "0\"x\"", "Assignment: project a 10x larger version."),
        ("Chinchilla tokens per parameter", 20, BLUE, "0", "Source: Hoffmann et al., 'Training Compute-Optimal "
         "Large Language Models', arXiv:2203.15556 (2022): ~20 tokens/param."),
        ("T4 peak dense FP16 (TFLOPS)", 65, BLUE, "0", "Source: NVIDIA T4 datasheet, 65 TFLOPS mixed-precision "
         "(FP16/FP32) Tensor Core peak."),
        ("A100 peak dense BF16 (TFLOPS)", 312, BLUE, "0", "Source: NVIDIA A100 datasheet, 312 TFLOPS BF16/FP16 "
         "dense Tensor Core peak."),
        ("Model FLOPs utilization (MFU)", 0.3, BLUE, "0%", "Assumption: 30% of peak is typical for small "
         "transformers; tiny models often achieve less."),
        ("T4 price ($/GPU-hour)", 0.35, BLUE, "$0.00", "Assumption: typical on-demand cloud T4 rate. Colab free "
         "tier is $0 but session-capped."),
        ("A100 price ($/GPU-hour)", 1.50, BLUE, "$0.00", "Assumption: typical on-demand cloud A100 40GB rate."),
        ("Mean tokens per request (test set)", avg_tokens, BLUE, "0.0", "From data/dataset_card.json: mean bytes "
         "per test example + 1 CLS token."),
    ]
    refs = {}
    for i, (label, val, color, fmt, note) in enumerate(inputs, start=5):
        put(ws, f"A{i}", label)
        put(ws, f"B{i}", val, color, fmt, fill=INPUT_FILL if color == BLUE else None, note=note)
        refs[label] = f"$B${i}"
    N, D, U = refs["Baseline parameters N"], refs["Baseline training tokens D (as trained)"], \
        refs["Unique training tokens per epoch"]
    S, K = refs["Scale factor"], refs["Chinchilla tokens per parameter"]
    T4, A100, MFU = refs["T4 peak dense FP16 (TFLOPS)"], refs["A100 peak dense BF16 (TFLOPS)"], \
        refs["Model FLOPs utilization (MFU)"]
    T4P, A100P, TOK = refs["T4 price ($/GPU-hour)"], refs["A100 price ($/GPU-hour)"], \
        refs["Mean tokens per request (test set)"]

    top = 5 + len(inputs) + 1
    header_row(ws, top, ["Projection", "Baseline (as trained)", "Baseline, Chinchilla-optimal", "10x params, same data",
                         "10x params, Chinchilla-optimal"])
    ws.row_dimensions[top].height = 32
    r0 = top + 1
    rows = [
        ("Parameters N", [f"={N}", f"={N}", f"={N}*{S}", f"={N}*{S}"], "#,##0"),
        ("Training tokens D", [f"={D}", f"={K}*C{r0}", f"={D}", f"={K}*E{r0}"], "#,##0"),
        ("Tokens per parameter (D/N)", [f"={c}{r0 + 1}/{c}{r0}" for c in "BCDE"], "0.0"),
        ("Training FLOPs (6ND)", [f"=6*{c}{r0}*{c}{r0 + 1}" for c in "BCDE"], "0.00E+00"),
        ("Passes over our current unique data", [f"={c}{r0 + 1}/{U}" for c in "BCDE"], "#,##0.0"),
        ("T4 GPU-hours at MFU", [f"={c}{r0 + 3}/({T4}*1E12*{MFU})/3600" for c in "BCDE"], "#,##0.00"),
        ("T4 cost ($)", [f"={c}{r0 + 5}*{T4P}" for c in "BCDE"], "$#,##0.00"),
        ("A100 GPU-hours at MFU", [f"={c}{r0 + 3}/({A100}*1E12*{MFU})/3600" for c in "BCDE"], "#,##0.00"),
        ("A100 cost ($)", [f"={c}{r0 + 7}*{A100P}" for c in "BCDE"], "$#,##0.00"),
        ("Weights at FP32 (MB)", [f"={c}{r0}*4/2^20" for c in "BCDE"], "#,##0.0"),
        ("Weights at 4-bit (MB)", [f"={c}{r0}*0.5/2^20" for c in "BCDE"], "#,##0.0"),
        ("Inference GFLOPs per request (~2N per token)", [f"=2*{c}{r0}*{TOK}/1E9" for c in "BCDE"], "0.000"),
    ]
    for i, (label, formulas, fmt) in enumerate(rows):
        r = r0 + i
        put(ws, f"A{r}", label, bold=label.startswith("Training FLOPs"))
        for col, f in zip("BCDE", formulas):
            put(ws, f"{col}{r}", f, fmt=fmt, bold=label.startswith("Training FLOPs"))
        if i % 2:
            for col in "ABCDE":
                ws[f"{col}{r}"].fill = BAND_FILL
    end = r0 + len(rows) - 1

    h = end + 2
    put(ws, f"A{h}", "Headline for the board", bold=True)
    ws[f"A{h + 1}"] = (f'="A "&TEXT({S},"0")&"x model ("&TEXT(E{r0},"#,##0")&" params) trained compute-optimally '
                       f'needs "&TEXT(E{r0 + 1},"#,##0")&" tokens = "&TEXT(E{r0 + 3},"0.0E+00")&" FLOPs, about "'
                       f'&TEXT(E{r0 + 5},"#,##0.0")&" T4-hours ($"&TEXT(E{r0 + 6},"#,##0")&"). That is "'
                       f'&TEXT(E{r0 + 4},"#,##0")&" passes over the data we have today, so new labeled data, '
                       f'not GPU time, is the binding cost."')
    ws[f"A{h + 1}"].font = font(bold=True)
    ws[f"A{h + 1}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(f"A{h + 1}:E{h + 3}")
    ws[f"A{h + 5}"] = ("Notes: 6ND counts forward + backward matmul FLOPs and ignores attention's sequence-length "
                       "term (small at our ~150-token mean). 2N per token similarly understates inference cost for "
                       "long documents. Chinchilla's 20 tokens/param was fit on large language models; it is used "
                       "here as the standard rule of thumb, not a law for 3M-parameter classifiers.")
    ws[f"A{h + 5}"].font = font(GREY, italic=True)
    ws[f"A{h + 5}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(f"A{h + 5}:E{h + 8}")
    set_widths(ws, {"A": 44, "B": 22, "C": 24, "D": 22, "E": 26})


def build_recommendation(wb, board_first, board_last, ranking):
    ws = wb.create_sheet("Recommendation")
    title(ws, "Ranked recommendation and residual risks",
          "Rank 1 is the configuration to ship. Edit results/recommendation.json and rebuild to change it.")
    header_row(ws, 4, ["Rank", "Config", "Accuracy", "p95 ms", "Shipped MB", "Why (one line for Sandra)",
                       "Residual risks she accepts"])
    dv = DataValidation(type="list", formula1=f"=Board!$A${board_first}:$A${board_last}", allow_blank=True)
    ws.add_data_validation(dv)
    for i in range(len(NAMES)):
        r = 5 + i
        put(ws, f"A{r}", i + 1, bold=True)
        entry = ranking[i] if i < len(ranking) else {}
        put(ws, f"B{r}", entry.get("config"), BLUE, fill=INPUT_FILL)
        dv.add(f"B{r}")
        rng = f"Board!$A${board_first}:$A${board_last}"
        for col, src, fmt in [("C", "F", "0.0%"), ("D", "J", "0.00"), ("E", "M", "0.00")]:
            put(ws, f"{col}{r}", f"=IF(B{r}=\"\",\"\",INDEX(Board!${src}${board_first}:${src}${board_last},"
                                 f"MATCH(B{r},{rng},0)))", GREEN, fmt)
        put(ws, f"F{r}", entry.get("why"), BLUE, fill=INPUT_FILL)
        put(ws, f"G{r}", entry.get("risks"), BLUE, fill=INPUT_FILL)
        ws.row_dimensions[r].height = 66 if entry else 15
        for col in "FG":
            ws[f"{col}{r}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["B5"].comment = Comment("Pick a config from the list; accuracy / latency / size fill in from the Board.",
                               "benchmark")

    top = 5 + len(NAMES) + 2
    put(ws, f"A{top}", "Residual-risk register (draft - edit, re-score, or delete rows)", bold=True)
    header_row(ws, top + 1, ["#", "Risk", "Applies to", "Likelihood", "Impact", "Mitigation"])
    risks = [
        ("Latency was measured on a Colab T4 GPU, not the in-cab Android tablet. Absolute ms will be higher, and "
         "the ranking can change (4-bit and linear-attention kernels are tuned for GPUs).", "all",
         "High", "Medium", "On-device pilot on the target tablet (ONNX Runtime / ExecuTorch) before fleet rollout."),
        ("Document accuracy comes from synthetic scans. Real bills of lading add handwriting, stamps, multi-page "
         "scans and layouts we did not simulate.", "all", "High", "High",
         "Label ~500 real scans from the document archive and re-run the benchmark before sign-off."),
        ("Driver notes are scored on clean SLURP transcripts. On-device speech-to-text errors and cab noise are "
         "not modeled, and note accuracy is already well below document accuracy.", "all", "High", "Medium",
         "Benchmark against real ASR output; keep a driver confirm/correct step in the app."),
        ("4-bit here uses bitsandbytes NF4, which is CUDA-only. The tablet needs a different int4 runtime whose "
         "accuracy and speed have not been measured.", "quant_4bit", "Medium", "Medium",
         "Re-quantize with the tablet's runtime and re-check accuracy on this test set."),
        ("Every model was trained once with one seed; training-run variance is not measured (only inference "
         "trials are repeated).", "all", "Medium", "Low", "Retrain the shortlisted config with 3 seeds."),
        ("No reject option: a document type outside the 8 classes is forced into one of them.", "all", "Medium",
         "Medium", "Add a confidence threshold that routes low-confidence scans to manual review."),
    ]
    for i, (risk, applies, like, impact, mit) in enumerate(risks, start=1):
        r = top + 1 + i
        put(ws, f"A{r}", i)
        for col, val in zip("BCDEF", (risk, applies, like, impact, mit)):
            put(ws, f"{col}{r}", val, BLUE)
            ws[f"{col}{r}"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 54
    lv = DataValidation(type="list", formula1='"Low,Medium,High"', allow_blank=True)
    ws.add_data_validation(lv)
    lv.add(f"D{top + 2}:E{top + 1 + len(risks) + 10}")
    set_widths(ws, {"A": 6, "B": 60, "C": 14, "D": 11, "E": 11, "F": 50, "G": 50})
    ws.column_dimensions["B"].width = 60


def build_protocol(wb, card):
    ws = wb.create_sheet("Protocol")
    title(ws, "Measurement protocol", "Identical for every configuration (scripts/benchmark.py).")
    rows = [
        ("Hardware", f'=IF(COUNTA(Runs!$A:$A)<2,"(no runs yet)",{first_run("device")})'),
        ("Software", f'=IF(COUNTA(Runs!$A:$A)<2,"(no runs yet)","torch "&{first_run("torch_version")}&", CUDA "'
                     f'&{first_run("cuda_version")}&", Python "&{first_run("python_version")})'),
        ("Seeds", "Fixed seed 627 for data generation, training, and request order (trial k uses seed 627+k)."),
        ("Request shape", "Batch size 1: one scanned document or one spoken note per request, as on the tablet."),
        ("Latency", "Host tensor -> GPU -> forward -> argmax back on host, bracketed by torch.cuda.synchronize(); "
                    "p50 / p95 over every request in a trial."),
        ("Warm-up", "50 untimed requests before every trial (CUDA context, kernel selection, caches)."),
        ("Repeated trials", "5 full passes over the test set per configuration; Board shows the mean across "
                            "trials and the p50 spread."),
        ("Peak memory", "torch.cuda.max_memory_allocated() during the timed trial (weights + activations)."),
        ("Shipped size", "Serialized state_dict size as loaded; 4-bit stores packed NF4 weights + scales."),
        ("Accuracy", "From the same timed pass, so accuracy and latency describe the same predictions."),
        ("Precision", "Baseline, linear attention and student run FP32; 4-bit runs NF4 weights with FP16 matmuls."),
    ]
    for i, (k, v) in enumerate(rows, start=4):
        put(ws, f"A{i}", k, bold=True)
        put(ws, f"B{i}", v, GREEN if str(v).startswith("=") else BLACK)
        ws[f"B{i}"].alignment = Alignment(wrap_text=True, vertical="top")

    r = 4 + len(rows) + 1
    put(ws, f"A{r}", "Evaluation set", bold=True)
    header_row(ws, r + 1, ["Split", "Source", "Examples", "Mean bytes", "p95 bytes"])
    rr = r + 2
    for split, info in (card.get("splits") or {}).items():
        for src, s in info["by_source"].items():
            for col, val in zip("ABCDE", (split, src, s["n"], s["bytes_mean"], s["bytes_p95"])):
                put(ws, f"{col}{rr}", val)
            rr += 1
    ws[f"A{rr + 1}"] = ("synthetic_docs: scripts/generate_documents.py, 8 freight document types with OCR noise, "
                        "cut-off headers and partial photos; the test split uses unseen company/city names.  "
                        "slurp_notes: SLURP transcripts (github.com/pswietojanski/slurp, CC BY 4.0), 8 in-cab "
                        "intents, SLURP's own splits.")
    ws[f"A{rr + 1}"].font = font(GREY, italic=True)
    set_widths(ws, {"A": 18, "B": 90, "C": 12, "D": 12, "E": 12})


def main():
    runs = load_runs()
    training = load_training()
    card_path = DATA_DIR / "dataset_card.json"
    card = json.loads(card_path.read_text()) if card_path.exists() else {}
    test = (card.get("splits") or {}).get("test", {}).get("by_source", {})
    n_test = sum(s["n"] for s in test.values())
    avg_tokens = round(sum(s["n"] * s["bytes_mean"] for s in test.values()) / n_test + 1, 1) if n_test else 150.0

    wb = Workbook()
    last = build_runs(wb, runs)  # created first so other sheets can reference its size
    ranking = json.loads(RECOMMENDATION.read_text())["ranking"] if RECOMMENDATION.exists() else []
    unknown = [e["config"] for e in ranking if e["config"] not in CONFIGS]
    if unknown:
        sys.exit(f"{RECOMMENDATION} names unknown configs: {unknown}")
    first, lastrow = build_board(wb, runs, last, ranking)
    build_training(wb, training)
    build_scaling(wb, avg_tokens)
    build_recommendation(wb, first, lastrow, ranking)
    build_protocol(wb, card)
    wb.move_sheet("Board", offset=-wb.index(wb["Board"]))

    RESULTS_DIR.mkdir(exist_ok=True)
    wb.save(WORKBOOK)
    configs_run = sorted({r["config"] for r in runs})
    print(f"{len(runs)} trial rows from {len(configs_run)} configs {configs_run} -> {WORKBOOK}")


if __name__ == "__main__":
    main()
